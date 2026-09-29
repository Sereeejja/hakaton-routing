#!/usr/bin/env python3
"""Build the Routecraft submission deck from the official LCT 2026 template."""

from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


DARK = RGBColor(0x1C, 0x1D, 0x22)
MUTED = RGBColor(0x67, 0x69, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PINK = RGBColor(0xFF, 0x00, 0x53)
PURPLE = RGBColor(0x52, 0x09, 0x78)
LAVENDER = RGBColor(0x8A, 0x83, 0xD1)
ACID = RGBColor(0xD9, 0xFF, 0x43)
PALE = RGBColor(0xF3, 0xF4, 0xEF)


def set_text(shape, text: str, *, size: int | None = None, bold: bool | None = None,
             color: RGBColor | None = None, align=None) -> None:
    if not shape.has_text_frame:
        return
    shape.text = text
    frame = shape.text_frame
    frame.word_wrap = True
    for paragraph in frame.paragraphs:
        if align is not None:
            paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = "Montserrat"
            if size is not None:
                run.font.size = Pt(size)
            if bold is not None:
                run.font.bold = bold
            if color is not None:
                run.font.color.rgb = color


def add_text(slide, text: str, x: float, y: float, w: float, h: float, *,
             size: int = 16, bold: bool = False, color: RGBColor = DARK,
             align=PP_ALIGN.LEFT) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.margin_left = 0
    shape.text_frame.margin_right = 0
    shape.text_frame.margin_top = 0
    shape.text_frame.margin_bottom = 0
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(shape, text, size=size, bold=bold, color=color, align=align)


def add_round_rect(slide, x: float, y: float, w: float, h: float, fill: RGBColor,
                   *, radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE, line: RGBColor | None = None):
    shape = slide.shapes.add_shape(radius_shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def set_slide_number(slide, value: int) -> None:
    for shape in slide.shapes:
        if shape.has_text_frame and "Номер слайда" in shape.name:
            set_text(shape, str(value), size=12, bold=True)


def add_route_mock(slide, x: float, y: float, w: float, h: float) -> None:
    """Draw an editable product mock directly with PowerPoint shapes."""

    add_round_rect(slide, x, y, w, h, PALE, line=RGBColor(0xD9, 0xDC, 0xD4))
    sidebar = add_round_rect(slide, x, y, w * 0.29, h, DARK)
    sidebar.adjustments[0] = 0.05
    add_text(slide, "routecraft", x + 0.14, y + 0.12, w * 0.23, 0.28,
             size=10, bold=True, color=WHITE)
    add_text(slide, "Заявки  8\nБригады  3", x + 0.14, y + 0.62, w * 0.22, 0.58,
             size=7, color=RGBColor(0xC8, 0xCC, 0xC3))
    button = add_round_rect(slide, x + 0.14, y + h - 0.54, w * 0.21, 0.32, ACID)
    add_text(slide, "ПОСТРОИТЬ", x + 0.19, y + h - 0.50, w * 0.18, 0.22,
             size=6, bold=True, color=DARK, align=PP_ALIGN.CENTER)

    map_x = x + w * 0.33
    map_y = y + 0.15
    map_w = w * 0.62
    map_h = h - 0.3
    for offset in (0.16, 0.38, 0.61, 0.80):
        road = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(map_x), Inches(map_y + map_h * offset), Inches(map_w), Inches(0.025),
        )
        road.fill.solid(); road.fill.fore_color.rgb = RGBColor(0xD2, 0xD6, 0xCF)
        road.line.color.rgb = RGBColor(0xD2, 0xD6, 0xCF)
        road.rotation = -8 if int(offset * 100) % 2 else 9
    points = [
        (map_x + 0.15, map_y + map_h * 0.72),
        (map_x + map_w * 0.42, map_y + map_h * 0.53),
        (map_x + map_w * 0.60, map_y + map_h * 0.28),
        (map_x + map_w * 0.85, map_y + map_h * 0.42),
    ]
    for first, second in zip(points, points[1:]):
        line = slide.shapes.add_connector(1, Inches(first[0]), Inches(first[1]), Inches(second[0]), Inches(second[1]))
        line.line.color.rgb = PINK
        line.line.width = Pt(3)
    for index, (px, py) in enumerate(points):
        marker = add_round_rect(slide, px - 0.09, py - 0.09, 0.18, 0.18,
                                ACID if index == 0 else PINK,
                                radius_shape=MSO_SHAPE.OVAL, line=WHITE)
        marker.line.width = Pt(1.5)


def fill_mandatory(prs: Presentation) -> None:
    title, overview, team, story, short = [prs.slides[index] for index in range(6, 11)]

    set_text(title.shapes[0], "routecraft", size=36, bold=True)
    set_text(title.shapes[2], "Интеллектуальное планирование рабочих маршрутов\nдля выездных инженеров", size=17)
    beeline_logo = prs.slides[5].shapes[20].image.blob
    title.shapes.add_picture(BytesIO(beeline_logo), Inches(0.45), Inches(0.34), width=Inches(2.10))
    add_text(title, "ЛЦТ 2026  ·  ЗАДАЧА №3", 2.85, 0.45, 3.4, 0.35,
             size=10, bold=True, color=WHITE)
    add_text(title, "План на день — за один клик", 7.2, 5.95, 5.3, 0.55,
             size=22, bold=True, color=PURPLE, align=PP_ALIGN.RIGHT)

    set_text(overview.shapes[6], "ROUTECRAFT", size=25, bold=True)
    set_text(overview.shapes[4],
             "Капитан: Максим Егоров, разработчик\nКол-во участников: 2 человека\nКоманда: backend, оптимизация и продуктовая web-разработка\nГород и регион: Москва",
             size=12, color=WHITE)
    set_text(overview.shapes[2], "Краткое описание решения:", size=13, bold=True)
    set_text(overview.shapes[11], "Сервис распределяет заявки между бригадами и строит выполнимые маршруты по дорогам с учётом временных окон, навыков, транспорта и смен.", size=14)
    set_text(overview.shapes[3], "Диспетчер работает с картой, а технические параметры скрыты. Каждое назначение объясняется; авария перестраивает уже опубликованный план.", size=14)
    add_route_mock(overview, 0.20, 0.48, 6.55, 2.48)

    set_text(team.shapes[20], "КОМАНДА ROUTECRAFT", size=24, bold=True)
    set_text(team.shapes[22], "Максим Егоров", size=14, bold=True)
    set_text(team.shapes[21], "Product & frontend\nGitHub: @DefaultMaxim\nКонтакт: профиль команды ЛЦТ", size=10)
    set_text(team.shapes[2], "Сергей", size=14, bold=True)
    set_text(team.shapes[1], "Backend & optimization\nGitHub: @Sereeejja\nКонтакт: профиль команды ЛЦТ", size=10)
    for placeholder_index, initials, color in [(15, "МЕ", PINK), (16, "С", PURPLE)]:
        placeholder = team.shapes[placeholder_index]
        x, y, w, h = (placeholder.left / 914400, placeholder.top / 914400,
                      placeholder.width / 914400, placeholder.height / 914400)
        avatar = add_round_rect(team, x, y, w, h, color, radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(team, initials, x, y, w, h, size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # Keep the two left cards and remove the three unused template cards.
    for index in sorted([3, 4, 5, 6, 7, 8, 9, 10, 11, 17, 18, 19], reverse=True):
        remove_shape(team.shapes[index])

    set_text(story.shapes[2], "КАК МЫ СОБРАЛИ РЕШЕНИЕ", size=23, bold=True)
    set_text(story.shapes[4], "Краткая история команды", size=13, bold=True)
    set_text(story.shapes[3], "Два разработчика объединили компетенции в исследовании операций, backend и интерфейсах. Работали короткими итерациями: модель → API → карта → проверка сценария.", size=12)
    set_text(story.shapes[7], "Ручная диспетчеризация плохо масштабируется: нужно одновременно удерживать окна клиентов, квалификации, транспорт и срочные аварии.", size=12)
    set_text(story.shapes[8], "Почему выбрали задачу", size=13, bold=True)
    set_text(story.shapes[5], "Главные вызовы — неоднородные адреса, недоступная дорожная матрица и противоречивые даты. Добавили кэш, безопасный fallback, строгую валидацию и аудит источников.", size=12)
    set_text(story.shapes[6], "Сложности и как мы их преодолели", size=13, bold=True)

    set_text(short.shapes[6], "КОРОТКО О РЕШЕНИИ", size=23, bold=True)
    set_text(short.shapes[4],
             "Python-ядро решает VRPTW с открытыми маршрутами. Go REST API хранит заявки, бригады и версии планов в PostgreSQL. React-карта показывает реальные дорожные линии, последовательность остановок и анимацию движения. Перепланирование сохраняет предыдущие назначения, где это возможно.", size=13)
    set_text(short.shapes[5],
             "Routecraft превращает сложную оптимизацию в понятный рабочий инструмент диспетчера: точки ставятся кликом, план строится одной кнопкой, причины неназначения видны сразу. Решение готово к пилоту на одном офисе и расширению на сеть сервисных зон.", size=13)


def fill_details(prs: Presentation) -> list:
    problem = prs.slides[23]
    set_text(problem.shapes[15], "ПРОБЛЕМА → РЕШЕНИЕ → ЭФФЕКТ", size=22, bold=True)
    set_text(problem.shapes[7], "Диспетчер вручную сопоставляет десятки заявок, окна клиентов и разные компетенции. Авария ломает уже собранный день.", size=14)
    set_text(problem.shapes[8], "Routecraft автоматически назначает бригады, упорядочивает остановки и строит путь по дорожной сети.", size=14)
    set_text(problem.shapes[9], "Больше выполненных заявок, ни одной потерянной аварии, меньше ручных решений и прозрачные причины каждого назначения.", size=14)

    flow = prs.slides[24]
    set_text(flow.shapes[26], "КАК ЭТО РАБОТАЕТ", size=22, bold=True)
    stages = [
        (2, 3, "01 · ВХОД", "Заявки, офисы и мастер-данные бригад"),
        (4, 5, "02 · МАТРИЦА", "Время и геометрия пути по дорогам"),
        (6, 7, "03 · ПЛАН", "Окна, навыки, смены и приоритеты"),
        (8, 9, "04 · ДИСПЕТЧЕР", "Карта, объяснения и анимация"),
        (10, 11, "05 · СОБЫТИЕ", "Авария или недоступность → новый план"),
    ]
    for title_idx, body_idx, heading, body in stages:
        set_text(flow.shapes[title_idx], heading, size=12, bold=True)
        set_text(flow.shapes[body_idx], body, size=10)

    architecture = prs.slides[19]
    set_text(architecture.shapes[14], "АРХИТЕКТУРА", size=22, bold=True)
    set_text(architecture.shapes[7],
             "React + TypeScript\n\nИнтерактивная карта диспетчера\nПостановка точек кликом\nДорожные маршруты и анимация\nОбъяснение назначений",
             size=15, bold=False, color=WHITE)
    architecture_items = [
        "Go REST API · валидация и сценарии перепланирования",
        "PostgreSQL · заявки, бригады, планы, события и миграции",
        "Python solver · OR-Tools VRPTW + baseline + HGS",
        "OSRM · матрица времени и дорожная геометрия",
        "Swagger + тесты · воспроизводимый контракт и приёмка",
    ]
    for idx, value in zip([8, 9, 10, 11, 12], architecture_items):
        set_text(architecture.shapes[idx], value, size=12, bold=True)

    constraints = prs.slides[17]
    set_text(constraints.shapes[14], "ЧТО УЧИТЫВАЕТ ПЛАН", size=22, bold=True)
    pairs = [
        (2, 3, "01", "Временные окна\nначало визита — внутри обещанного интервала"),
        (4, 5, "02", "Навыки\nподключение, ремонт и аварийные работы"),
        (6, 7, "03", "График\n2/2 или 5/2 и индивидуальные границы смены"),
        (8, 9, "04", "Транспорт\nавто, пешком, велосипед, общественный"),
        (10, 11, "05", "Приоритет\nавария → подключение → остальные"),
        (12, 13, "06", "Открытый маршрут\nстарт из офиса, возврат не обязателен"),
    ]
    for title_idx, body_idx, num, body in pairs:
        set_text(constraints.shapes[title_idx], num, size=14, bold=True, color=PINK)
        set_text(constraints.shapes[body_idx], body, size=11)

    results = prs.slides[20]
    set_text(results.shapes[12], "ПРОВЕРКА НА СИНТЕТИЧЕСКОМ СЦЕНАРИИ", size=20, bold=True)
    chart_data = CategoryChartData()
    chart_data.categories = ["Baseline", "Routecraft"]
    chart_data.add_series("Выполнено заявок", (13, 27))
    results.shapes[11].chart.replace_data(chart_data)
    set_text(results.shapes[5], "27 из 30", size=24, bold=True, color=PINK)
    set_text(results.shapes[6], "выполнено против 13 у прозрачного baseline", size=11)
    set_text(results.shapes[7], "0", size=24, bold=True, color=PINK)
    set_text(results.shapes[8], "неназначенных срочных заявок; у baseline — 3", size=11)
    set_text(results.shapes[9], "3 бригады", size=24, bold=True, color=PINK)
    set_text(results.shapes[10], "задействовано вместо 4 при большей полноте", size=11)
    add_text(results, "Сценарий: 30 заявок, 4 бригады. Метрики не являются производственным прогнозом.",
             0.55, 6.75, 5.9, 0.28, size=7, color=MUTED)

    demo = prs.slides[25]
    set_text(demo.shapes[7], "ДЕМО ДИСПЕТЧЕРА", size=22, bold=True)
    add_route_mock(demo, 4.45, 1.26, 4.42, 3.45)
    set_text(demo.shapes[4], "1 · ТОЧКИ\nЗаявка или старт бригады добавляются кликом по карте.", size=11)
    set_text(demo.shapes[5], "2 · МАРШРУТ\nЛинии идут по дорогам; старт, остановки и финиш различимы.", size=11)
    set_text(demo.shapes[6], "3 · ИЗМЕНЕНИЕ\nОтмена или недоступность бригады сразу перестраивают план.", size=11)

    data = prs.slides[16]
    set_text(data.shapes[10], "ДАННЫЕ И ПРОВЕРКА КАЧЕСТВА", size=22, bold=True)
    data_cards = [
        (4, 5, 11, "3 ЗОНЫ", "Независимые офисы и наборы бригад; назначения контроля не подмешиваются."),
        (6, 7, 12, "497 ЗАЯВОК", "Дополнительные дни проверены; 19 аварий найдены по типу HD."),
        (8, 9, 13, "5 КОНТУРОВ", "Go, Python, React, TypeScript и production build проходят одной командой."),
    ]
    for title_idx, body_idx, num_idx, heading, body in data_cards:
        set_text(data.shapes[num_idx], heading.split()[0], size=17, bold=True, color=PINK)
        set_text(data.shapes[title_idx], heading, size=14, bold=True)
        set_text(data.shapes[body_idx], body, size=11)

    roadmap = prs.slides[14]
    set_text(roadmap.shapes[0], "ПЛАН ВНЕДРЕНИЯ", size=22, bold=True)
    roadmap_items = [
        "1 · Пилот на одном офисе: реальные мастер-данные и KPI диспетчера",
        "2 · Локальный OSRM и промышленный геокодер с корпоративным SLA",
        "3 · Производственный календарь и многодневный горизонт планирования",
        "4 · Телеметрия: факт прибытия, длительность работ, качество прогнозов",
        "5 · Масштабирование на сервисные зоны и интеграция с BeKeeper",
    ]
    for index, value in zip(range(1, 6), roadmap_items):
        set_text(roadmap.shapes[index], value, size=13, bold=index == 1)

    final = prs.slides[12]
    set_text(final.shapes[3], "ROUTECRAFT ГОТОВ К ПИЛОТУ", size=22, bold=True)
    set_text(final.shapes[1],
             "Публичный репозиторий\nhttps://github.com/Sereeejja/hakaton-routing\n\nВетка решения\nalgo_routing\n\nЛокальная проверка\n./scripts/verify.sh\n\nПродукт: карта диспетчера\nAPI: /swagger/index.html\n\nКонтакт команды — в профиле ЛЦТ",
             size=17, color=WHITE)
    add_text(final, "План на день — за один клик.", 7.35, 5.65, 5.25, 0.6,
             size=24, bold=True, color=PURPLE, align=PP_ALIGN.RIGHT)

    return [problem, flow, architecture, constraints, results, demo, data, roadmap, final]


def select_and_order_slides(prs: Presentation, detail_slides: list) -> None:
    mandatory = [prs.slides[index] for index in range(6, 11)]
    desired = mandatory + detail_slides
    desired_ids = [slide.slide_id for slide in desired]
    id_to_element = {int(item.id): item for item in list(prs.slides._sldIdLst)}
    for item in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(item)
    for slide_id in desired_ids:
        prs.slides._sldIdLst.append(id_to_element[slide_id])
    for number, slide in enumerate(list(prs.slides)[5:], start=12):
        set_slide_number(slide, number)


def validate_deck(prs: Presentation) -> None:
    forbidden = (
        "Опишите в чем", "Имя Фамилия", "Название команды", "Расскажите, как вы",
        "Что делает ваше решение", "В чем суть вашего решения",
    )
    all_text = "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    leftovers = [value for value in forbidden if value.casefold() in all_text.casefold()]
    if leftovers:
        raise RuntimeError(f"Template placeholders remain: {leftovers}")
    if len(prs.slides) != 14:
        raise RuntimeError(f"Expected 14 slides, got {len(prs.slides)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("template", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    prs = Presentation(args.template)
    fill_mandatory(prs)
    detail_slides = fill_details(prs)
    select_and_order_slides(prs, detail_slides)
    validate_deck(prs)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print(f"saved {len(prs.slides)} slides to {args.output}")


if __name__ == "__main__":
    main()
